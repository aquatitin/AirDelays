#!/usr/bin/env python3
"""Ingesta del conjunt «Airport Arrival ATFM Delay» d'Eurocontrol.

Font: EUROCONTROL Aviation Intelligence Portal (ansperformance.eu/data/).
Un registre per aeroport i dia amb les arribades IFR i els minuts de
retard ATFM (Air Traffic Flow Management) imputats a l'aeroport
d'arribada. Cobreix els aeroports de l'àrea ECAC des del 2014 i
s'actualitza mensualment.

L'script baixa el fitxer complet (CSV o XLSX segons la variant
disponible), agrega per mes a nivell d'aeroport i d'estat, i escriu
docs/data/eurocontrol/.

Ús:
    python3 pipeline/eurocontrol_atfm.py            # baixa, agrega i escriu
    python3 pipeline/eurocontrol_atfm.py --probe    # només mostra URL i capçaleres
    python3 pipeline/eurocontrol_atfm.py --file Airport_Arrival_ATFM_Delay.xlsx
"""

from __future__ import annotations

import argparse
import bz2
import csv
import io
import re
import sys
import os
import urllib.error

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from datastore import Dataset, http_get  # noqa: E402

# Variants conegudes de publicació del fitxer; es prova en ordre.
CANDIDATE_URLS = [
    "https://www.eurocontrol.int/performance/data/download/csv/Airport_Arrival_ATFM_Delay.csv.bz2",
    "https://www.eurocontrol.int/performance/data/download/csv/Airport_Arrival_ATFM_Delay.csv",
    "https://ansperformance.eu/download/csv/Airport_Arrival_ATFM_Delay.csv.bz2",
    "https://ansperformance.eu/download/csv/Airport_Arrival_ATFM_Delay.csv",
    "https://www.eurocontrol.int/performance/data/download/xls/Airport_Arrival_ATFM_Delay.xlsx",
    "https://ansperformance.eu/download/xls/Airport_Arrival_ATFM_Delay.xlsx",
]

META = {
    "order": 1,
    "name": "Europa — retard ATFM per aeroport (Eurocontrol)",
    "url": "https://ansperformance.eu/data/",
    "attribution": "Font: EUROCONTROL (ansperformance.eu)",
    "rec": "ec",
    "dims": ["airport", "country"],
    "note": (
        "Retard de regulació de trànsit aeri (ATFM) imputat a l'aeroport "
        "d'arribada. No inclou retards propis de les aerolínies (rotació, "
        "tripulacions, tècnics): és l'indicador de congestió d'aeroport i "
        "espai aeri que publica Eurocontrol."
    ),
}

# rec "ec": [arribades, minuts_retard_atfm]
REC_LEN = 2


def fetch(path_or_none: str | None) -> tuple[str, bytes]:
    """Retorna (nom, contingut). Prova les URL candidates en ordre."""
    if path_or_none:
        with open(path_or_none, "rb") as f:
            return os.path.basename(path_or_none), f.read()
    last: Exception | None = None
    for url in CANDIDATE_URLS:
        try:
            print(f"Provant {url}…", flush=True)
            data = http_get(url, retries=1)
            # Les pàgines d'error retornen HTML; el fitxer bo és gran.
            if len(data) < 10_000 and b"<html" in data[:1000].lower():
                print("  resposta HTML inesperada; es descarta")
                continue
            print(f"  OK ({len(data) / 1e6:.1f} MB)")
            return url.rsplit("/", 1)[-1], data
        except urllib.error.HTTPError as exc:
            print(f"  HTTP {exc.code}")
            last = exc
        except Exception as exc:  # noqa: BLE001
            print(f"  error: {exc}")
            last = exc
    raise SystemExit(f"Cap URL candidata ha funcionat (últim error: {last})")


def rows_from_payload(name: str, data: bytes):
    """Itera diccionaris de fila des de CSV, CSV.BZ2 o XLSX."""
    lower = name.lower()
    if lower.endswith(".bz2"):
        data = bz2.decompress(data)
        lower = lower[:-4]
    if lower.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        yield from csv.DictReader(io.StringIO(text), delimiter=delimiter)
        return
    if lower.endswith((".xlsx", ".xls")):
        import openpyxl  # dependència només per a la variant Excel

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = max(wb.worksheets, key=lambda w: w.max_row or 0)
        rows = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
        for row in rows:
            yield dict(zip(header, row))
        return
    raise SystemExit(f"Format no reconegut: {name}")


def col(row: dict, *names: str):
    """Cerca un valor per nom de columna, sense distingir majúscules."""
    for name in names:
        if name in row:
            return row[name]
    lowered = {str(k).strip().lower(): v for k, v in row.items()}
    for name in names:
        if name.lower() in lowered:
            return lowered[name.lower()]
    return None


def to_int(value) -> int:
    if value is None or value == "":
        return 0
    return int(round(float(value)))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", help="fitxer local ja baixat (opcional)")
    parser.add_argument("--probe", action="store_true", help="mostra capçaleres i surt")
    args = parser.parse_args()

    name, data = fetch(args.file)
    rows = rows_from_payload(name, data)

    first = next(rows, None)
    if first is None:
        raise SystemExit("El fitxer no conté files")
    print("Columnes detectades:", list(first.keys()))
    if args.probe:
        print("Primera fila:", first)
        return

    ds = Dataset("eurocontrol", META, REC_LEN)
    n = 0
    skipped = 0

    def ingest(row: dict) -> None:
        nonlocal n, skipped
        apt = col(row, "APT_ICAO")
        state = col(row, "STATE_NAME")
        year = col(row, "YEAR")
        month = col(row, "MONTH_NUM")
        arrivals = col(row, "FLT_ARR_1")
        delay = col(row, "DLY_APT_ARR_1")
        if year in (None, "") or month in (None, ""):
            # Variant sense YEAR/MONTH_NUM: es deriva de FLT_DATE (ISO 8601).
            flt_date = str(col(row, "FLT_DATE") or "")
            match = re.match(r"(\d{4})-(\d{2})", flt_date)
            if match:
                year, month = match.group(1), match.group(2)
        if not apt or year in (None, "") or month in (None, ""):
            skipped += 1
            return
        year, month = int(float(year)), int(float(month))
        values = [to_int(arrivals), to_int(delay)]
        ds.add(year, month, "airport", str(apt).strip(), values)
        apt_name = col(row, "APT_NAME")
        if apt_name:
            ds.set_name("airport", str(apt).strip(), str(apt_name).strip())
        if state:
            state = str(state).strip()
            ds.add(year, month, "country", state, values)
            ds.set_name("country", state, state.title() if state.isupper() else state)
        n += 1

    ingest(first)
    for row in rows:
        ingest(row)

    ds.write()
    print(f"Fet: {n} files agregades ({skipped} omeses) -> docs/data/eurocontrol/")


if __name__ == "__main__":
    main()
